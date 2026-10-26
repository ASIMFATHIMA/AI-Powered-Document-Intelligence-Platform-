
from io import BytesIO
import fitz
from docx import Document


def extract_text(uploaded_file):
    filename = uploaded_file.name.lower()
    content = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        with fitz.open(stream=content, filetype="pdf") as pdf:
            text = "\n".join(page.get_text() for page in pdf)

    elif filename.endswith(".docx"):
        doc = Document(BytesIO(content))

        paragraphs = [
            paragraph.text
            for paragraph in doc.paragraphs
            if paragraph.text.strip()
        ]

        table_text = []
        for table in doc.tables:
            for row in table.rows:
                table_text.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        text = "\n".join(paragraphs + table_text)

    elif filename.endswith(".txt"):
        text = content.decode("utf-8", errors="replace")

    else:
        raise ValueError("Unsupported file format.")

    if not text.strip():
        raise ValueError(
            "No text found. Scanned PDFs require OCR."
        )

    return text


def chunk_text(text, chunk_size=4000, overlap=300):
    chunks = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunks.append(text[start:end])

        if end == len(text):
            break

        start = end - overlap

    return chunks
